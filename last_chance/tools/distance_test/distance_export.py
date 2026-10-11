"""Excel, CSV, and three-curve plots for real or explicitly simulated sessions."""
import csv
import hashlib
import json
import os
from datetime import datetime, timezone
from pathlib import Path
from distance_core import SENSORS, build_tables, load_events


def export_session(directory, plots=True):
    directory = Path(directory)
    meta = json.loads((directory / "session.json").read_text(encoding="utf-8"))
    events = load_events(directory)
    samples, captures, summary = build_tables(events, meta)
    # Capture rows were built with original settings for legacy attempts.
    # The workbook overview shows the latest explicitly appended settings.
    for event in events:
        if event["kind"] == "capture_settings":
            meta.update({key: event[key] for key in ("samples_per_capture", "period_ms", "settle_ms")})
    open_ended = meta.get("collection_mode") == "open_ended" or any(
        e["kind"] == "collection_mode" and e["mode"] == "open_ended" for e in events)
    if not samples:
        print("No samples yet; session metadata and event logs are saved.")
        return None
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S_%fZ")
    output = directory / "analysis" / stamp
    output.mkdir(parents=True)
    datasets = {"samples": samples, "captures": captures, "summary": summary}
    for name, rows in datasets.items():
        if not rows:
            continue
        with (output / f"{name}.csv").open("w", newline="", encoding="utf-8-sig") as stream:
            writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)
    hashes = {p.name: hashlib.sha256(p.read_bytes()).hexdigest()
              for p in [directory / "session.json", *sorted(directory.glob("events-*.jsonl"))]}
    (output / "sources.json").write_text(json.dumps(hashes, indent=2), encoding="utf-8")
    sensor = SENSORS[meta["sensor"]]
    title = ("SIMULATED - " if meta.get("simulated") else "") + sensor["name"]
    quantities = [("raw", f"Sensor output ({sensor['raw_unit']})")]
    if meta["sensor"] == "ir":
        quantities.append(("voltage", f"Output voltage (V; {meta['adc_reference_source']} {meta['adc_reference_v']:g} V reference)"))
    else:
        quantities.append(("calibrated_cm", "Reported distance (cm, from sensor face)"))
    workbook_available = write_workbook(output, datasets, meta, title, quantities)
    if plots and summary:
        os.environ.setdefault("MPLCONFIGDIR", str((output / "plot-cache").resolve()))
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        figure, axes = plt.subplots(1, 2, figsize=(12, 4.6), constrained_layout=True)
        for ax, (quantity, label) in zip(axes, quantities):
            x = [r["actual_cm"] for r in summary]
            for statistic, color in (("min", "#2878b5"), ("mean", "#24823d"), ("max", "#ce472f")):
                y = [r[f"{quantity}_{statistic}"] if r[f"{quantity}_{statistic}"] is not None else float("nan") for r in summary]
                ax.plot(x, y, "o-", label=statistic.capitalize(), color=color, markersize=4)
            if meta["sensor"] == "ir" and max(x, default=0) > 80:
                ax.axvspan(80, max(x), color="0.7", alpha=0.2, label="Outside 10-80 cm specification")
            ax.set(xlabel="Measured sensor-to-target distance (cm)", ylabel=label)
            ax.grid(alpha=0.25)
            ax.legend(fontsize=8)
        figure.suptitle(title + " - minimum / mean / maximum of capture means")
        for extension in ("png", "svg"):
            figure.savefig(output / f"characteristic.{extension}", dpi=180)
        plt.close(figure)
    coverage = (f"{len(summary)} measured positions (open-ended session)." if open_ended
                else f"{len(meta['points_cm'])} planned.")
    lines = [f"# {title}", "", f"Samples: {len(samples)}; capture attempts: {len(captures)}.",
             f"Points with at least three usable accepted captures: {sum(r['three_usable_repeats'] for r in summary)}/{coverage}",
             "", "The horizontal axis uses the entered tape measurement from the sensor face.",
             "Each repetition is a separate capture. The three curves are the minimum, arithmetic mean,",
             "and maximum of accepted capture means at that position, with equal weight per capture.",
             "Sample-level spread is also retained in captures.csv. Missing echoes/ranges are not zero distances.",
             "Skipped, rejected, interrupted and superseded captures remain in samples.csv with inclusion flags.",
             "A point with fewer than three usable captures does not meet the repetition requirement.",
             "Live preview frames stay in the serial log; only explicit DATA captures enter the measurement tables.",
             "", "[Raw samples](samples.csv) | [Captures](captures.csv) | [Summary](summary.csv)"]
    if workbook_available:
        lines += ["", "[Excel workbook with charts](measurements.xlsx)"]
    if plots and summary:
        lines += ["", "![Characteristic](characteristic.png)"]
    if meta["sensor"] == "us":
        cfg = meta["firmware_config"]
        lines += ["", f"Ultrasonic conversion: raw_cm = echo_us / 58; calibrated_cm = raw_cm * {cfg['US_SCALE']} + {cfg['US_OFFSET_CM']}.",
                  "These existing coefficients were obtained over 15-150 cm. Outside that interval their use is extrapolation.",
                  "No robot-center mounting offset is added. Raw echo values and uncorrected cm are retained."]
    elif meta["sensor"] == "ir":
        lines += ["", "The Sharp GP2Y0A21YK0F specification covers 10-80 cm. Farther points are exploratory ADC observations.",
                  "There is no measured IR distance calibration in the project; no generic inverse-distance formula is applied.",
                  f"Volts = ADC * {meta['adc_reference_v']} / 1023; reference source: {meta['adc_reference_source']}."]
    else:
        lines += ["", f"Recorded laser profile: {meta['firmware_config'].get('LASER_PROFILE', 'not recorded')}. Output is mm; there is no extra empirical laser correction.",
                  "Successful device status 11 is required. The library does not implement ST's full host sigma validation.",
                  "The 2 m endpoint is an exploration of actual range; target reflectance, lighting and profile affect success."]
    lines += ["", "This characterizes the tested setup. It does not establish the sensor's entire specified range or absolute accuracy.",
              "Existing calibration constants are preserved; this program does not apply new coefficients."]
    (output / "report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"Results: {output.resolve()}")
    return output


def write_workbook(output, datasets, meta, title, quantities):
    try:
        from openpyxl import Workbook
        from openpyxl.chart import Reference, ScatterChart, Series
    except ImportError:
        print("Excel workbook requires openpyxl; CSV files were saved. Install requirements.txt and run analyze again.")
        return False
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Experiment"
    sheet.append([title])
    sheet.append(["Curve definition", "Minimum / mean / maximum of separate accepted capture means"])
    for key, value in meta.items():
        sheet.append([key, json.dumps(value) if isinstance(value, (dict, list)) else value])
    for name, rows in datasets.items():
        ws = workbook.create_sheet(name.title())
        if not rows:
            continue
        columns = list(rows[0])
        ws.append(columns)
        for row in rows:
            ws.append([row[k] for k in columns])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        if name == "summary":
            for index, (quantity, label) in enumerate(quantities):
                chart = ScatterChart()
                chart.title = title
                chart.x_axis.title = "Measured distance (cm)"
                chart.y_axis.title = label
                chart.style = 13
                chart.width, chart.height = 22, 12
                x = Reference(ws, min_col=columns.index("actual_cm") + 1, min_row=2, max_row=len(rows) + 1)
                for stat in ("min", "mean", "max"):
                    y = Reference(ws, min_col=columns.index(f"{quantity}_{stat}") + 1, min_row=2, max_row=len(rows) + 1)
                    chart.series.append(Series(y, x, title=stat.capitalize()))
                ws.add_chart(chart, f"A{len(rows)+4+index*25}")
    for ws in workbook:
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # notes are text, including leading '='
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 28
    workbook.save(output / "measurements.xlsx")
    return True
