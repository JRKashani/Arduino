import json
import io
import tempfile
import unittest
from argparse import Namespace
from pathlib import Path
from unittest.mock import patch
from contextlib import redirect_stdout

from distance_core import build_tables, convert, finite_number, load_events, parse_sample
from distance_mission import Device, LiveConsole, live_text, record
from distance_export import export_session


def metadata(sensor="us"):
    return {"sensor": sensor, "firmware_config": {"US_SCALE": "1.015350", "US_OFFSET_CM": "0.709100"},
            "adc_reference_v": 5.0, "adc_reference_source": "nominal", "points_cm": [100], "repeats": 3}


def point(pid="a", actual=100):
    return {"kind": "point", "point_id": pid, "point_index": 1, "target_cm": 100, "actual_cm": actual}


def capture(cid, raw_values, pid="a", decision="accepted", valid=True, complete=True):
    events = [{"kind": "capture", "capture_id": cid, "point_id": pid, "repeat": cid}]
    events += [{"kind": "sample", "capture_id": cid, "sequence": i, "timestamp_ms": i*100,
                "raw": raw, "valid": valid, "status": "OK" if valid else "NO_ECHO", "device_status": -1}
               for i, raw in enumerate(raw_values)]
    if complete:
        events.append({"kind": "capture_complete", "capture_id": cid})
    if decision:
        events.append({"kind": "review", "capture_id": cid, "decision": decision})
    return events


class SilentLog:
    def __init__(self):
        self.events = []

    def write(self, kind, **values):
        self.events.append({"kind": kind, **values})

    def wire(self, direction, line):
        pass


class FakePort:
    def __init__(self, lines):
        self.lines = iter(lines)
        self.sent = []

    def write(self, data):
        self.sent.append(data)

    def flush(self):
        pass

    def read_until(self, delimiter):
        return next(self.lines, b"")


class RecordingDevice:
    def __init__(self, port, log):
        self.log = log

    def connect(self):
        return "test", metadata()["firmware_config"]

    def send(self, command):
        pass

    def enable_monitor(self, sensor):
        pass

    def capture(self, cid, sensor, count, period, settle):
        rows = []
        for sequence in range(count):
            row = {"capture_id": cid, "sequence": sequence, "timestamp_ms": sequence*100,
                   "raw": 5800, "valid": True, "status": "OK", "device_status": -1}
            self.log.write("sample", **row)
            rows.append(row)
        self.log.write("capture_complete", capture_id=cid)
        return rows


class DistanceTests(unittest.TestCase):
    def run_recorder(self, directory, answers, resume=False):
        args = Namespace(resume=str(directory) if resume else None,
                         output=None if resume else str(directory), sensor="us", port="TEST",
                         repeats=3, samples=1, no_plots=True)
        with patch("serial.Serial"), patch("distance_mission.Device", RecordingDevice), \
             patch("distance_mission.LiveConsole.input", side_effect=answers), patch("distance_mission.export_session"), \
             patch("builtins.print"):
            record(args)

    def test_open_ended_recording_stops_after_user_chosen_distances(self):
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp) / "session"
            answers = ["", "198.7", *["", "a"]*3, "y", "63.2", *["", "a"]*3, "n"]
            self.run_recorder(directory, answers)
            meta = json.loads((directory / "session.json").read_text())
            events = load_events(directory)
            points = [e for e in events if e["kind"] == "point"]
            self.assertNotIn("points_cm", meta)
            self.assertEqual([p["actual_cm"] for p in points], [198.7, 63.2])
            self.assertEqual([p["point_index"] for p in points], [1, 2])
            self.assertEqual(sum(e["kind"] == "capture" for e in events), 6)
            self.assertEqual(meta["samples_per_capture"], 10)
            self.assertEqual(meta["period_ms"], 500)
            self.assertEqual(meta["settle_ms"] + 9*meta["period_ms"], 5000)
            self.assertEqual(events[-1]["kind"], "run_complete")

    def test_legacy_session_resume_allows_arbitrary_next_distance(self):
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            meta = {**metadata(), "points_cm": [200, 175, 150], "samples_per_capture": 1,
                    "period_ms": 100, "settle_ms": 1000}
            (directory / "session.json").write_text(json.dumps(meta), encoding="utf-8")
            original = [point(actual=198.7)] + capture(1, [5800]) + [{"kind": "point_done", "point_id": "a"}]
            source = directory / "events-000.jsonl"
            source.write_text("\n".join(json.dumps(e) for e in original)+"\n", encoding="utf-8")
            original_bytes = source.read_bytes()
            self.run_recorder(directory, ["y", "62.8", *["", "a"]*3, "n"], resume=True)
            self.assertEqual(source.read_bytes(), original_bytes)
            points = [e for e in load_events(directory) if e["kind"] == "point"]
            self.assertEqual([p["actual_cm"] for p in points], [198.7, 62.8])
            self.assertEqual(points[-1]["target_cm"], 62.8)
            _, captures, _ = build_tables(load_events(directory), meta)
            self.assertEqual(captures[0]["period_ms"], 100)
            self.assertTrue(all(c["period_ms"] == 500 for c in captures[1:]))

    def test_unfinished_point_resume_restarts_and_preserves_attempts(self):
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            meta = {**metadata(), "collection_mode": "open_ended", "samples_per_capture": 1,
                    "period_ms": 100, "settle_ms": 1000}
            (directory / "session.json").write_text(json.dumps(meta), encoding="utf-8")
            events = [point(actual=100)] + capture(1, [9000])
            (directory / "events-000.jsonl").write_text("\n".join(json.dumps(e) for e in events)+"\n", encoding="utf-8")
            self.run_recorder(directory, ["99.4", *["", "a"]*3, "n"], resume=True)
            samples, _, summary = build_tables(load_events(directory), meta)
            self.assertFalse(samples[0]["included"])
            self.assertEqual(summary[0]["actual_cm"], 99.4)
            self.assertEqual(summary[0]["accepted_captures"], 3)

    def test_existing_ultrasonic_calibration_has_no_robot_center_offset(self):
        result = convert({"valid": True, "raw": 5800}, metadata())
        self.assertEqual(result["raw_cm"], 100)
        self.assertAlmostEqual(result["calibrated_cm"], 102.2441)

    def test_no_echo_retained_but_never_zero_distance(self):
        events = [point()] + capture(1, [0, 0], valid=False)
        samples, captures, summary = build_tables(events, metadata())
        self.assertEqual(samples[0]["raw"], 0)
        self.assertIsNone(samples[0]["calibrated_cm"])
        self.assertIsNone(summary[0]["raw_mean"])
        self.assertEqual(summary[0]["usable_captures"], 0)

    def test_three_curves_are_equally_weighted_capture_means(self):
        events = [point()] + capture(1, [1000]*10) + capture(2, [2000]) + capture(3, [3000]*2)
        samples, captures, summary = build_tables(events, metadata())
        row = summary[0]
        self.assertEqual((row["raw_min"], row["raw_mean"], row["raw_max"]), (1000, 2000, 3000))
        self.assertTrue(row["three_usable_repeats"])

    def test_rejected_and_partial_attempts_stay_in_samples_not_curves(self):
        events = [point()] + capture(1, [1000]) + capture(2, [9000], decision="rejected")
        events += capture(3, [8000], complete=False, decision=None)
        samples, _, summary = build_tables(events, metadata())
        self.assertEqual(len(samples), 3)
        self.assertEqual([r["included"] for r in samples], [True, False, False])
        self.assertEqual(summary[0]["raw_mean"], 1000)

    def test_remeasured_point_supersedes_previous_partial_point(self):
        events = [point()] + capture(1, [1000]) + [point("b", 99.4)] + capture(2, [2000], pid="b")
        samples, _, summary = build_tables(events, metadata())
        self.assertFalse(samples[0]["included"])
        self.assertEqual(summary[0]["actual_cm"], 99.4)
        self.assertEqual(summary[0]["raw_mean"], 2000)

    def test_skip_on_resume_excludes_previous_partial_point(self):
        events = [point()] + capture(1, [1000]) + [{"kind": "point_skipped", "point_index": 1}]
        samples, _, summary = build_tables(events, metadata())
        self.assertFalse(samples[0]["included"])
        self.assertEqual(summary, [])

    def test_ir_out_of_range_and_rail_still_record_real_adc(self):
        events = [point(actual=200)] + capture(1, [0, 10])
        samples, _, summary = build_tables(events, metadata("ir"))
        self.assertFalse(summary[0]["within_nominal_range"])
        self.assertEqual(summary[0]["raw_mean"], 5)
        self.assertIsNone(summary[0]["calibrated_cm_mean"])
        self.assertAlmostEqual(samples[1]["voltage"], 50/1023)

    def test_ultrasonic_200cm_is_extrapolated(self):
        _, _, summary = build_tables([point(actual=200)] + capture(1, [11600]), metadata())
        self.assertTrue(summary[0]["within_nominal_range"])
        self.assertFalse(summary[0]["within_us_calibration_range"])

    def test_protocol_rejects_lost_frames_and_invalid_validity(self):
        for line in ("DATA,1,1,100,5800,1,OK,-1", "DATA,2,0,100,5800,1,OK,-1", "DATA,1,0,100,0,1,OK,-1"):
            with self.assertRaises(ValueError):
                parse_sample(line, 1, 0, "us")
        sample = parse_sample("DATA,1,0,100,11600,1,OK,-1", 1, 0, "us")
        self.assertEqual(sample["raw"], 11600)

    def test_invalid_laser_status_cannot_be_used(self):
        with self.assertRaises(ValueError):
            parse_sample("DATA,1,0,100,2000,1,OK,4", 1, 0, "laser")
        sample = parse_sample("DATA,1,0,100,8190,0,RANGE_STATUS,4", 1, 0, "laser")
        self.assertFalse(sample["valid"])

    def test_partial_serial_lines_and_reset_detection(self):
        log = SilentLog()
        device = Device(FakePort([b"BEGIN,1,U,1\n", b"DATA,1,0,100,58", b"00,1,OK,-1\n", b"END,1,1,complete\n"]), log)
        samples = device.capture(1, "us", 1, 100, 1000)
        self.assertEqual(samples[0]["raw"], 5800)
        self.assertEqual(log.events[-1]["kind"], "capture_complete")
        device = Device(FakePort([b"BEGIN,1,U,1\n", b"BOOT,DISTANCE,1\n"]), SilentLog())
        with self.assertRaisesRegex(RuntimeError, "reset"):
            device.capture(1, "us", 1, 100, 1000)

    def test_nonfinite_and_negative_distance_rejected(self):
        for value in ("nan", "inf", "-1", "0"):
            with self.assertRaises(ValueError):
                finite_number(value, 1, 400)

    def test_preview_does_not_enter_captured_sample_data(self):
        log = SilentLog()
        device = Device(FakePort([b"BEGIN,1,U,1\n", b"LIVE,U,100,5800,1,OK,-1,cm=100,corrected_cm=102.24\n",
                                  b"DATA,1,0,100,5800,1,OK,-1\n", b"END,1,1,complete\n"]), log)
        displayed = []
        device.live_callback = displayed.append
        rows = device.capture(1, "us", 1, 500, 500)
        self.assertEqual(len(rows), 1)
        self.assertEqual(sum(e["kind"] == "sample" for e in log.events), 1)
        self.assertEqual(len(displayed), 1)

    def test_live_monitor_selection_and_formatting(self):
        device = Device(FakePort([b"MONITOR,U,500\n"]), SilentLog())
        device.enable_monitor("us")
        self.assertEqual(device.port.sent, [b"MONITOR,U\n"])
        text = live_text("LIVE,U,100,5800,1,OK,-1,cm=100", metadata())
        self.assertIn("corrected=102.24 cm", text)
        self.assertIn("NO_ECHO", live_text("LIVE,U,100,0,0,NO_ECHO,-1", metadata()))
        self.assertIsNone(live_text("LIVE,I,100,500,1,OK,-1", metadata()))

    def test_live_console_preserves_prompt_and_partially_typed_distance(self):
        console = LiveConsole(None, metadata())
        console.prompt, console.typed = "Distance: ", "198."
        output = io.StringIO()
        with redirect_stdout(output):
            console.display("LIVE,U,100,5800,1,OK,-1")
        self.assertIn("LIVE US", output.getvalue())
        self.assertTrue(output.getvalue().endswith("Distance: 198."))

    def test_export_contains_samples_and_three_excel_curves(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            (directory / "session.json").write_text(json.dumps(metadata()), encoding="utf-8")
            events = [point()] + capture(1, [5800]) + capture(2, [5900]) + capture(3, [6000])
            (directory / "events-test.jsonl").write_text("\n".join(json.dumps(e) for e in events)+"\n", encoding="utf-8")
            output = export_session(directory, plots=False)
            workbook = load_workbook(output / "measurements.xlsx")
            self.assertEqual(workbook["Samples"].max_row, 4)
            self.assertEqual(len(workbook["Summary"]._charts), 2)
            self.assertEqual(len(workbook["Summary"]._charts[0].series), 3)
            workbook.close()

    def test_open_ended_export_counts_entered_positions(self):
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            meta = metadata()
            meta.pop("points_cm")
            meta["collection_mode"] = "open_ended"
            (directory / "session.json").write_text(json.dumps(meta), encoding="utf-8")
            events = [point()] + capture(1, [5800]) + capture(2, [5900]) + capture(3, [6000])
            (directory / "events-test.jsonl").write_text("\n".join(json.dumps(e) for e in events)+"\n", encoding="utf-8")
            output = export_session(directory, plots=False)
            self.assertIn("1/1 measured positions (open-ended session)", (output / "report.md").read_text())


if __name__ == "__main__":
    unittest.main()
